"""Build editable Word files of docs/dossier and zip them per phase.

  pip install python-docx
  python scripts/dossier_naar_word.py      -> dist/vitascribe-dossier-<date>.zip

The zip holds intern/ (what phase 1 needs) and extern/ (everything). Only the
Markdown the dossier uses is supported: headings, paragraphs, lists, tables,
quotes, rules, and bold/italic/code inline.
"""

from __future__ import annotations

import datetime
import io
import pathlib
import re
import zipfile
from typing import List

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOSSIER = ROOT / "docs" / "dossier"
ALLEEN_EXTERN = {"11", "12"}
GROEN = RGBColor(0x04, 0x78, 0x57)
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def inline(par, tekst: str) -> None:
    # Links between the documents mean nothing inside Word; keep the label.
    tekst = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", tekst)
    for deel in INLINE.split(tekst):
        if not deel:
            continue
        if deel.startswith("**"):
            par.add_run(deel[2:-2]).bold = True
        elif deel.startswith("`"):
            par.add_run(deel[1:-1])
        elif deel.startswith("*") and len(deel) > 2:
            par.add_run(deel[1:-1]).italic = True
        else:
            par.add_run(deel)


def cellen(regel: str) -> List[str]:
    return [c.strip() for c in regel.strip().strip("|").split("|")]


def naar_docx(md: pathlib.Path) -> bytes:
    doc = Document()
    stijl = doc.styles["Normal"]
    stijl.font.name, stijl.font.size = "Calibri", Pt(11)
    regels = md.read_text(encoding="utf-8").splitlines()
    i, alinea = 0, []

    def schrijf_alinea() -> None:
        if alinea:
            inline(doc.add_paragraph(), " ".join(alinea))
            alinea.clear()

    while i < len(regels):
        r = regels[i]
        s = r.strip()
        lijst = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", r)
        if not s:
            schrijf_alinea()
        elif s.startswith("#"):
            schrijf_alinea()
            niveau = len(s) - len(s.lstrip("#"))
            kop = doc.add_heading(level=min(niveau, 3))
            inline(kop, s.lstrip("#").strip())
            for run in kop.runs:
                run.font.color.rgb = GROEN
        elif s.startswith("|"):
            schrijf_alinea()
            rijen = []
            while i < len(regels) and regels[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", regels[i].strip()):
                    rijen.append(cellen(regels[i]))
                i += 1
            breedte = max(len(x) for x in rijen)
            tabel = doc.add_table(rows=len(rijen), cols=breedte)
            tabel.style = "Table Grid"
            for y, rij in enumerate(rijen):
                for x in range(breedte):
                    par = tabel.cell(y, x).paragraphs[0]
                    inline(par, rij[x] if x < len(rij) else "")
                    if y == 0:
                        for run in par.runs:
                            run.bold = True
            continue
        elif s.startswith(">"):
            schrijf_alinea()
            blok = []
            while i < len(regels) and regels[i].strip().startswith(">"):
                blok.append(regels[i].strip()[1:].strip())
                i += 1
            for stuk in "\n".join(blok).split("\n\n"):
                for sub in stuk.split("\n- "):
                    tekst = sub.replace("\n", " ").strip()
                    if tekst.startswith("- "):
                        tekst = tekst[2:]
                    stijlnaam = "List Bullet" if sub is not stuk.split("\n- ")[0] or stuk.startswith("- ") else "Quote"
                    inline(doc.add_paragraph(style=stijlnaam), tekst)
            continue
        elif s == "---":
            schrijf_alinea()
            doc.add_paragraph("―" * 20)
        elif lijst:
            schrijf_alinea()
            tekst = lijst.group(3)
            while i + 1 < len(regels) and regels[i + 1].startswith("  ") and not re.match(r"^\s*(-|\d+\.)\s", regels[i + 1]):
                i += 1
                tekst += " " + regels[i].strip()
            diep = len(lijst.group(1)) >= 2
            if lijst.group(2)[0].isdigit():
                # Word's "List Number" keeps counting across lists; write the number.
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Pt(36 if diep else 18)
                par.paragraph_format.first_line_indent = Pt(-14)
                inline(par, lijst.group(2) + " " + tekst)
            else:
                inline(doc.add_paragraph(style="List Bullet" + (" 2" if diep else "")), tekst)
        else:
            alinea.append(s)
        i += 1
    schrijf_alinea()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def main() -> None:
    doel = ROOT / "dist" / f"vitascribe-dossier-{datetime.date.today().isoformat()}.zip"
    doel.parent.mkdir(exist_ok=True)
    with zipfile.ZipFile(doel, "w", zipfile.ZIP_DEFLATED) as z:
        for md in sorted(DOSSIER.glob("*.md")):
            naam = ("00-overzicht" if md.stem == "README" else md.stem) + ".docx"
            data = naar_docx(md)
            z.writestr(f"extern/{naam}", data)
            if naam[:2] not in ALLEEN_EXTERN:
                z.writestr(f"intern/{naam}", data)
    print(doel)


if __name__ == "__main__":
    main()
